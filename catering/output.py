"""
catering/output.py — Printable HTML checklist and Excel export.
"""
import datetime
import html

from catering.store import META_FIELDS

NAVY = "#1F4E78"   # print colour — reads well on paper and in black & white

PRINT_TEMPLATE = """<!doctype html>
<html><head><meta charset="utf-8"><title>__TITLE__</title>
<style>
  @page { size: letter; margin: 0.55in; }
  * { box-sizing: border-box; }
  body { font-family: "Segoe UI", Arial, sans-serif; color: #111; font-size: 10.5pt; margin: 0; }
  .page { max-width: 7.4in; margin: 0 auto; }
  header { border-bottom: 3px solid __NAVY__; padding-bottom: 8px; margin-bottom: 14px; }
  .kicker { font-size: 8.5pt; letter-spacing: .12em; text-transform: uppercase; color: #555; }
  h1 { font-size: 20pt; margin: 2px 0 10px; }
  .meta { display: flex; gap: 18px; }
  .meta div { flex: 1; border-bottom: 1px solid #999; padding: 2px 0; min-height: 22px; }
  .meta span { font-size: 8pt; text-transform: uppercase; color: #555; margin-right: 6px; }
  table { width: 100%; border-collapse: collapse; margin-bottom: 14px; }
  thead { display: table-header-group; }
  tr { page-break-inside: avoid; }
  tr.cat th { background: __NAVY__; color: #fff; text-align: left; padding: 5px 8px; font-size: 11pt;
              border: 1.5px solid #222; -webkit-print-color-adjust: exact; print-color-adjust: exact; }
  tr.cols th { font-size: 7.5pt; text-transform: uppercase; color: #333; text-align: left;
               padding: 4px 8px; border: 1.5px solid #222; background: #dde4ee;
               -webkit-print-color-adjust: exact; print-color-adjust: exact; }
  td { padding: 5px 8px; height: 28px; border: 1px solid #555; vertical-align: middle; }
  tbody tr:nth-child(even) td { background: #f4f6f9; -webkit-print-color-adjust: exact; print-color-adjust: exact; }
  tbody tr:last-child td { border-bottom: 1.5px solid #222; }
  tbody td:first-child { border-left: 1.5px solid #222; }
  tbody td:last-child { border-right: 1.5px solid #222; }
  .ck { width: 58px; text-align: center; }
  th.ck { text-align: center; }
  .qty { width: 95px; text-align: center; }
  th.qty { text-align: center; }
  .notes { width: 32%; color: #333; }
  .box { display: inline-block; width: 15px; height: 15px; border: 1.5px solid #222; border-radius: 2px;
         line-height: 12px; font-size: 12px; font-weight: bold; text-align: center; }
  .signoff { display: flex; gap: 18px; margin-top: 26px; page-break-inside: avoid; }
  .signoff div { flex: 1; border-top: 1px solid #333; padding-top: 3px; font-size: 8pt; color: #555;
                 text-transform: uppercase; margin-top: 28px; }
  footer { margin-top: 18px; font-size: 8pt; color: #777; display: flex; justify-content: space-between; }
  .bar { position: sticky; top: 0; background: #fff; padding: 10px 0; text-align: right; }
  .bar button { font: inherit; padding: 6px 16px; cursor: pointer; }
  @media screen { body { background: #e9ecef; } .page { background: #fff; padding: 0.5in; margin: 16px auto;
                  box-shadow: 0 1px 6px rgba(0,0,0,.2); } }
  @media print { .noprint { display: none; } }
</style></head>
<body onload="setTimeout(function(){ window.print(); }, 300)">
<div class="page">
  <div class="bar noprint"><button onclick="window.print()">Print</button></div>
  <header>
    <div class="kicker">J. Michael's Prime &middot; Catering Equipment Checklist</div>
    <h1>__TITLE__</h1>
    <div class="meta">__META__</div>
  </header>
  __SECTIONS__
  <div class="signoff"><div>Packed by</div><div>Loaded by</div><div>Returned &amp; checked in by</div></div>
  <footer><span>__COUNT__</span><span>Printed __PRINTED__</span></footer>
</div>
</body></html>
"""


def build_print_html(lst):
    e = html.escape

    def box(on):
        return f'<span class="box">{"&#10003;" if on else ""}</span>'

    meta = "".join(f"<div><span>{label}</span>{e(lst.get(key, ''))}</div>" for key, label in META_FIELDS)
    sections = []
    for cat in dict.fromkeys(i["category"] for i in lst["items"]):
        rows = "".join(
            f'<tr><td class="ck">{box(it["packed"])}</td><td class="ck">{box(it["loaded"])}</td>'
            f'<td>{e(it["item"])}</td><td class="qty">{e(it["qty"])}</td>'
            f'<td class="notes">{e(it["notes"])}</td></tr>'
            for it in lst["items"] if it["category"] == cat)
        sections.append(
            f'<table><thead><tr class="cat"><th colspan="5">{e(cat)}</th></tr>'
            '<tr class="cols"><th class="ck">Packed</th><th class="ck">Loaded</th><th>Item</th>'
            f'<th class="qty">Qty Needed</th><th>Notes</th></tr></thead><tbody>{rows}</tbody></table>')
    n = len(lst["items"])
    printed = datetime.datetime.now().strftime("%b %d, %Y %I:%M %p")
    return (PRINT_TEMPLATE.replace("__NAVY__", NAVY)
            .replace("__TITLE__", e(lst["name"]))
            .replace("__META__", meta)
            .replace("__SECTIONS__", "\n".join(sections) or "<p>No items on this checklist.</p>")
            .replace("__COUNT__", f"{n} item{'s' if n != 1 else ''}")
            .replace("__PRINTED__", printed))


def write_xlsx(lst, path):
    """Write one checklist to an .xlsx file. Raises ImportError / PermissionError."""
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill, Alignment, Border, Side

    wb = Workbook()
    ws = wb.active
    ws.title = "Checklist"
    ws.append([lst["name"]])
    ws["A1"].font = Font(bold=True, size=14)
    meta = "     ".join(f"{label}: {lst[key]}" for key, label in META_FIELDS if lst.get(key))
    if meta:
        ws.append([meta])
    ws.append([])
    ws.append(["Category", "Item", "Qty Needed", "Packed", "Loaded", "Notes"])
    hdr = ws.max_row
    for c in ws[hdr]:
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = PatternFill("solid", fgColor=NAVY[1:])
    for it in lst["items"]:
        ws.append([it["category"], it["item"], it["qty"], "✓" if it["packed"] else "",
                   "✓" if it["loaded"] else "", it["notes"]])
    for col, w in zip("ABCDEF", (26, 30, 14, 9, 9, 30)):
        ws.column_dimensions[col].width = w
    for row in ws.iter_rows(min_row=hdr + 1, min_col=3, max_col=5):
        for c in row:
            c.alignment = Alignment(horizontal="center")
    thin = Side(style="thin", color="555555")
    grid = Border(left=thin, right=thin, top=thin, bottom=thin)
    for row in ws.iter_rows(min_row=hdr, max_col=6):
        for c in row:
            c.border = grid
    wb.save(path)
